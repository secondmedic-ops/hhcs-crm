import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

MASTER_PATH = r'C:\Users\DELL\Documents\hhcs-crm\All_Projects_Work_and_Deployment_Log.xlsx'

from generate_today_work_report_2026_10_21 import TODAY_DATA

wb = openpyxl.load_workbook(MASTER_PATH)

font_data = Font(name="Calibri", size=11)
fill_even = PatternFill(fill_type="solid", start_color="FFFFFF")
fill_odd = PatternFill(fill_type="solid", start_color="F9FAFB")
border_thin = Side(border_style="thin", color="D9D9D9")
cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
align_left = Alignment(horizontal="left", vertical="center")
align_center = Alignment(horizontal="center", vertical="center")

# 1. Update 'Work Log (Requested Format)'
ws_work = wb['Work Log (Requested Format)']
# Insert rows at row 2
ws_work.insert_rows(2, amount=len(TODAY_DATA))

for idx, item in enumerate(TODAY_DATA, start=2):
    row_vals = [
        item["changes"],
        item["assigned_date"],
        item["completion_date"],
        item["start_time"],
        item["done_time"],
        item["status"],
        item["project"]
    ]
    ws_work.row_dimensions[idx].height = 24
    row_fill = fill_odd if idx % 2 == 1 else fill_even
    for col_idx, val in enumerate(row_vals, start=1):
        cell = ws_work.cell(row=idx, column=col_idx, value=val)
        cell.font = font_data
        cell.fill = row_fill
        cell.border = cell_border
        cell.alignment = align_center if col_idx in (2, 3, 4, 5, 6) else align_left

ws_work.auto_filter.ref = f"A1:G{ws_work.max_row}"

# 2. Update 'HHCS_CareDesk_CRM' sheet
if 'HHCS_CareDesk_CRM' in wb.sheetnames:
    ws_crm = wb['HHCS_CareDesk_CRM']
    crm_items = [item for item in TODAY_DATA if item["project"] == "HHCS CareDesk CRM"]
    ws_crm.insert_rows(2, amount=len(crm_items))
    for idx, item in enumerate(crm_items, start=2):
        row_vals = [
            item["changes"],
            item["assigned_date"],
            item["completion_date"],
            item["start_time"],
            item["done_time"],
            item["status"],
            item["project"]
        ]
        ws_crm.row_dimensions[idx].height = 24
        row_fill = fill_odd if idx % 2 == 1 else fill_even
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws_crm.cell(row=idx, column=col_idx, value=val)
            cell.font = font_data
            cell.fill = row_fill
            cell.border = cell_border
            cell.alignment = align_center if col_idx in (2, 3, 4, 5, 6) else align_left
    ws_crm.auto_filter.ref = f"A1:G{ws_crm.max_row}"

# 3. Update 'SM-ERP' sheet
if 'SM-ERP' in wb.sheetnames:
    ws_erp = wb['SM-ERP']
    erp_items = [item for item in TODAY_DATA if item["project"] == "SM-ERP"]
    ws_erp.insert_rows(2, amount=len(erp_items))
    for idx, item in enumerate(erp_items, start=2):
        row_vals = [
            item["changes"],
            item["assigned_date"],
            item["completion_date"],
            item["start_time"],
            item["done_time"],
            item["status"],
            item["project"]
        ]
        ws_erp.row_dimensions[idx].height = 24
        row_fill = fill_odd if idx % 2 == 1 else fill_even
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws_erp.cell(row=idx, column=col_idx, value=val)
            cell.font = font_data
            cell.fill = row_fill
            cell.border = cell_border
            cell.alignment = align_center if col_idx in (2, 3, 4, 5, 6) else align_left
    ws_erp.auto_filter.ref = f"A1:G{ws_erp.max_row}"

wb.save(MASTER_PATH)
print("Successfully updated master workbook:", MASTER_PATH)
